import os
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
import copy

# ==========================================================
# CONFIGURATION
# ==========================================================

COLOR_THEME = {
    "Participated": "#347363",
    "Not Participated": "#FFFFFF",
    "English": "#604A7B",
    "Mathematics": "#77933C",
    "Science": "#E46C0A",
}


def get_theme_color(label, default=None):
    hex_color = COLOR_THEME.get(label, default)
    return hex_color


SUBJECT_MAP = {
    "ENG": "English",
    "MATH": "Mathematics",
    "SCI": "Science",
    "EVS": "Environmental Science",
    "HIN": "Hindi",
}

TITLES = {
    "MASTERY": "Mastery Levels Analysis",
    "QLO": "Learning Outcome Performance",
    "SCHOOL": "Schoolwise Subject Performance",
}

CHART_DPI = 300
IMAGE_FOLDER = "temp_charts"

# ==========================================================
# STANDARD CHART FACTORY
# ==========================================================


def create_standard_chart(figsize=(10, 5)):
    fig, ax = plt.subplots(figsize=figsize)

    # Structural discipline
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis="y", linestyle="--", alpha=0.3)

    return fig, ax


# ==========================================================
# PATHS
# ==========================================================

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
INPUT_FILE = os.path.join(DATA_DIR, "uploadable_data.xlsx")
OUTPUT_FILE = os.path.join(DATA_DIR, "ANALYSIS_REPORT.pptx")

os.makedirs(IMAGE_FOLDER, exist_ok=True)

# ==========================================================
# UTILITIES
# ==========================================================


def normalize_columns(df):
    df.columns = df.columns.str.strip().str.replace(r"\s+", "_", regex=True).str.upper()
    return df


def apply_dynamic_ylim(ax, values):
    if len(values) == 0:
        return

    upper = float(max(values))

    if upper == 0:
        upper = 1  # prevent flat axis collapse

    ax.set_ylim(0, upper * 1.3)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for border_name in ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]:
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        tag_name = border_name.split(":")[1]
        existing = tcPr.find(f"{ns}{tag_name}")
        if existing is not None:
            tcPr.remove(existing)

    insert_pos = 0
    for border_name in ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]:
        ln = OxmlElement(border_name)
        ln.set("w", "12700")
        ln.set("cmpd", "sng")
        solidFill = OxmlElement("a:solidFill")
        srgbClr = OxmlElement("a:srgbClr")
        srgbClr.set("val", "000000")
        solidFill.append(srgbClr)
        ln.append(solidFill)

        prstDash = OxmlElement("a:prstDash")
        prstDash.set("val", "solid")
        ln.append(prstDash)

        tcPr.insert(insert_pos, ln)
        insert_pos += 1


def load_sheet(xls, name):
    try:
        return normalize_columns(pd.read_excel(xls, sheet_name=name))
    except Exception:
        return pd.DataFrame()


def duplicate_template_slide(prs, index=0):
    source = prs.slides[index]
    layout = source.slide_layout
    new_slide = prs.slides.add_slide(layout)

    for shape in source.shapes:
        if not shape.is_placeholder:
            new_el = copy.deepcopy(shape.element)
            new_slide.shapes._spTree.insert_element_before(new_el, "p:extLst")

    return new_slide


def enforce_template_title(slide, prs, text):

    green_bar = get_top_green_bar(prs.slides[0], prs)

    if not green_bar:
        return

    # remove existing textboxes inside green bar
    for shape in slide.shapes:
        if shape.has_text_frame:
            if abs(shape.top - green_bar.top) < Inches(0.1) and abs(
                shape.height - green_bar.height
            ) < Inches(0.1):
                slide.shapes._spTree.remove(shape.element)

    title_box = slide.shapes.add_textbox(
        green_bar.left, green_bar.top, green_bar.width, green_bar.height
    )

    tf = title_box.text_frame
    tf.clear()
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER

    for run in p.runs:
        run.font.name = "Fraunces 72pt"
        run.font.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(255, 255, 255)


def get_all_shapes(slide):
    shapes = []
    if getattr(slide, "slide_layout", None):
        if getattr(slide.slide_layout, "slide_master", None):
            shapes.extend(list(slide.slide_layout.slide_master.shapes))
        shapes.extend(list(slide.slide_layout.shapes))
    shapes.extend(list(slide.shapes))
    return shapes


def get_top_green_bar(slide, prs):

    bars = [
        s
        for s in get_all_shapes(slide)
        if hasattr(s, "width")
        and s.width > prs.slide_width * 0.6
        and hasattr(s, "height")
        and s.height < prs.slide_height * 0.25
        and hasattr(s, "top")
        and s.top < prs.slide_height * 0.5
    ]

    if not bars:
        return None

    return min(bars, key=lambda s: s.top)


def get_top_safe_y(slide, prs):
    green_bar = get_top_green_bar(prs.slides[0], prs)
    if green_bar:
        return int(green_bar.top + green_bar.height)
    return int(prs.slide_height * 0.18)


def get_bottom_safe_y(slide, prs):
    bars = [
        s
        for s in get_all_shapes(prs.slides[0])
        if hasattr(s, "top") and s.top > prs.slide_height * 0.6
    ]
    if not bars:
        return int(prs.slide_height * 0.85)
    return int(min(s.top for s in bars))


def add_centered_chart(slide, prs, image_path, reserve_bottom=0):

    slide_width = prs.slide_width

    top_safe_y = get_top_safe_y(slide, prs)
    bottom_safe_y = get_bottom_safe_y(slide, prs)

    usable_height = (bottom_safe_y - top_safe_y) - reserve_bottom

    pic_max_width = int(slide_width * 0.75)
    pic_max_height = int(max(1, usable_height * 0.95))

    pic = slide.shapes.add_picture(image_path, 0, 0)

    width_ratio = float(pic_max_width) / float(pic.width)
    height_ratio = float(pic_max_height) / float(pic.height)
    scale = min(width_ratio, height_ratio)

    final_width = int(pic.width * scale)
    final_height = int(pic.height * scale)

    pic.width = final_width
    pic.height = final_height
    pic.left = int((slide_width - final_width) / 2.0)
    pic.top = int(top_safe_y + (usable_height - final_height) / 2.0)


# ==========================================================
# MASTERY CHART
# ==========================================================


def generate_mastery_chart(grade_df, grade):

    subjects = sorted(grade_df["SUBJECT"].dropna().unique())
    levels = sorted(grade_df["QLVL"].dropna().unique())

    subject_labels = [SUBJECT_MAP.get(s, s) for s in subjects]

    fig, ax = create_standard_chart((10, 5))

    ax.text(
        0.5,
        1.05,
        f"Grade: {grade}",
        transform=ax.transAxes,
        ha="center",
        va="bottom",
        fontsize=12,
        fontweight="bold",
        color="#333333",
    )

    x = np.arange(len(subjects))
    total_levels = len(levels)
    bar_width = 0.8 / total_levels

    for i, lvl in enumerate(levels):
        values = []
        counts = []

        for subject in subjects:
            match = grade_df[
                (grade_df["SUBJECT"] == subject) & (grade_df["QLVL"] == lvl)
            ]

            if not match.empty:
                values.append(float(match["PERCENTAGE"].values[0]))
                counts.append(int(match["COUNT"].values[0]))
            else:
                values.append(0)
                counts.append(0)

        offsets = x - 0.4 + (i + 0.5) * bar_width
        bars = ax.bar(offsets, values, bar_width, edgecolor="black", linewidth=1)

        for j, bar in enumerate(bars):
            if values[j] == 0 and counts[j] == 0:
                continue

            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + 1,
                f"{counts[j]} ({values[j]:.0f}%)",
                ha="center",
                va="bottom",
                fontsize=8,
                fontweight="bold",
            )

            ax.text(
                bar.get_x() + bar.get_width() / 2,
                1,
                f"L{lvl}",
                ha="center",
                va="bottom",
                fontsize=7,
                fontweight="bold",
            )

    ax.set_xticks(x)
    ax.set_xticklabels(subject_labels)
    ax.set_ylim(0, 100)
    ax.set_ylabel("Percentage", labelpad=10, fontweight="bold", fontsize=12)
    ax.set_xlabel("Subjects", labelpad=10, fontweight="bold", fontsize=12)

    if ax.get_legend():
        ax.get_legend().remove()

    plt.tight_layout()

    image_path = os.path.join(IMAGE_FOLDER, f"grade_{grade}_mastery.png")
    plt.savefig(image_path, dpi=CHART_DPI)
    plt.close()

    return image_path


# ==========================================================
# QLO CHARTS
# ==========================================================


def generate_qlo_chart(sub_df, subject):
    df = sub_df.copy()
    df["AVERAGE_PERFORMANCE"] = pd.to_numeric(
        df["AVERAGE_PERFORMANCE"], errors="coerce"
    ).fillna(0)
    df = df.sort_values("AVERAGE_PERFORMANCE", ascending=True)

    fig, ax = create_standard_chart((10, max(4, len(df) * 0.6)))
    ax.grid(axis="y", visible=False)
    ax.grid(axis="x", linestyle="--", alpha=0.3)

    y = np.arange(len(df))
    height = 0.6

    theme_subj = SUBJECT_MAP.get(subject, subject)
    bar_color = get_theme_color(theme_subj, default="#5b9bd5")

    bars = ax.barh(
        y,
        df["AVERAGE_PERFORMANCE"],
        height=height,
        color=bar_color,
        edgecolor="black",
        linewidth=1,
    )

    for i, bar in enumerate(bars):
        width = bar.get_width()
        perf = int(round(width))
        questions = str(df.iloc[i]["QUESTIONS_LIST"])

        label_text = f"{perf}% ({questions})"

        ax.text(
            width + 2,
            bar.get_y() + bar.get_height() / 2,
            label_text,
            va="center",
            ha="left",
            fontweight="bold",
            fontsize=10,
        )

    wrapped_labels = [textwrap.fill(str(label), width=40) for label in df["QLO_FULL"]]

    ax.set_yticks(y)
    ax.set_yticklabels(wrapped_labels)
    ax.set_xlim(0, 100)
    ax.set_xlabel("Performance (%)", labelpad=10, fontweight="bold", fontsize=12)

    plt.tight_layout()

    safe_subj = str(subject).replace(" ", "_").replace("/", "_")
    image_path = os.path.join(IMAGE_FOLDER, f"qlo_{safe_subj}.png")
    plt.savefig(image_path, dpi=CHART_DPI, bbox_inches="tight")
    plt.close()

    return image_path


# ==========================================================
# SUBJECT PERFORMANCE CHARTS
# ==========================================================


def generate_overall_subject_performance(df):
    df = df.copy()
    df["SUB_ONLY"] = df["SUBJECT"].astype(str).str.split("-").str[-1].str.strip()

    summary = df.groupby("SUB_ONLY")["PERFORMANCE_PERCENTAGE"].mean().reset_index()

    fig, ax = create_standard_chart((10, 5))

    x = np.arange(len(summary))
    width = 0.6

    bar_colors = [get_theme_color(subj, default="blue") for subj in summary["SUB_ONLY"]]
    bars = ax.bar(
        x,
        summary["PERFORMANCE_PERCENTAGE"],
        width,
        color=bar_colors,
        edgecolor="black",
        linewidth=1,
    )

    for bar in bars:
        height = bar.get_height()
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            height + 1,
            f"{height:.0f}",
            ha="center",
            va="bottom",
            fontweight="bold",
            fontsize=10,
        )

    ax.set_xticks(x)
    ax.set_xticklabels(summary["SUB_ONLY"])
    ax.set_xlabel("Subjects", labelpad=10, fontweight="bold", fontsize=12)
    ax.set_ylabel("Performance (%)", labelpad=10, fontweight="bold", fontsize=12)
    ax.set_ylim(0, 100)

    plt.tight_layout()

    image_path = os.path.join(IMAGE_FOLDER, "overall_subject_performance.png")
    plt.savefig(image_path, dpi=CHART_DPI)
    plt.close()

    return image_path


def generate_per_school_subject_performance(df, school):
    df = df.copy()
    sub_df = df[df["SCHOOLNAME"] == school].copy()
    sub_df["SUB_ONLY"] = (
        sub_df["SUBJECT"].astype(str).str.split("-").str[-1].str.strip()
    )

    summary = sub_df.groupby("SUB_ONLY")["PERFORMANCE_PERCENTAGE"].mean().reset_index()

    fig, ax = create_standard_chart((10, 5))

    x = np.arange(len(summary))
    width = 0.6

    bar_colors = [get_theme_color(subj, default="blue") for subj in summary["SUB_ONLY"]]
    bars = ax.bar(
        x,
        summary["PERFORMANCE_PERCENTAGE"],
        width,
        color=bar_colors,
        edgecolor="black",
        linewidth=1,
    )

    for bar in bars:
        height = bar.get_height()
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            height + 1,
            f"{height:.0f}",
            ha="center",
            va="bottom",
            fontweight="bold",
            fontsize=10,
        )

    ax.set_xticks(x)
    ax.set_xticklabels(summary["SUB_ONLY"])
    ax.set_xlabel("Subjects", labelpad=10, fontweight="bold", fontsize=12)
    ax.set_ylabel("Performance (%)", labelpad=10, fontweight="bold", fontsize=12)
    ax.set_ylim(0, 100)

    plt.tight_layout()

    safe_name = str(school).replace(" ", "_").replace("/", "_")
    image_path = os.path.join(IMAGE_FOLDER, f"{safe_name}_subject_performance.png")
    plt.savefig(image_path, dpi=CHART_DPI)
    plt.close()

    return image_path


# ==========================================================
# PARTICIPATION CHARTS
# ==========================================================


def generate_overall_participation_pie(df):

    summary = df.groupby("TYPE")["COUNT"].sum()

    participated = summary.get("PARTICIPATED", 0)
    not_participated = summary.get("NOT_PARTICIPATED", 0)
    total_regs = participated + not_participated

    fig, ax = create_standard_chart((6, 5))
    ax.grid(False)

    def custom_autopct(pct, allvals):
        absolute = int(np.round(pct / 100.0 * np.sum(allvals)))
        return f"{absolute} ({pct:.0f}%)" if pct > 0 else ""

    colors = [get_theme_color("Participated"), get_theme_color("Not Participated")]

    wedges, texts, autotexts = ax.pie(
        [participated, not_participated],
        labels=["Participated", "Not Participated"],
        colors=colors,
        autopct=lambda pct: custom_autopct(pct, [participated, not_participated]),
        startangle=90,
        radius=0.75,
        center=(0, 0.4),
        wedgeprops=dict(edgecolor="black", linewidth=1),
    )

    ax.set_aspect("equal")

    for autotext in autotexts:
        autotext.set_fontsize(11)
        autotext.set_weight("bold")

    image_path = os.path.join(IMAGE_FOLDER, "overall_participation.png")
    plt.savefig(image_path, dpi=CHART_DPI)
    plt.close()

    return image_path, total_regs


def generate_school_participation_bar(df):

    pivot = df.pivot_table(
        index="SCHOOL_NAME", columns="TYPE", values="COUNT", fill_value=0
    ).reset_index()

    x = np.arange(len(pivot))
    width = 0.6

    part = pivot.get("PARTICIPATED", 0)
    not_part = pivot.get("NOT_PARTICIPATED", 0)

    fig, ax = create_standard_chart((10, 5))

    c_part = get_theme_color("Participated")
    c_npart = get_theme_color("Not Participated")

    p1 = ax.bar(
        x,
        part,
        width,
        label="Participated",
        color=c_part,
        edgecolor="black",
        linewidth=1,
    )
    p2 = ax.bar(
        x,
        not_part,
        width,
        bottom=part,
        label="Not Participated",
        color=c_npart,
        edgecolor="black",
        linewidth=1,
    )

    for bar in p1:
        if bar.get_height() > 0:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_y() + bar.get_height() / 2,
                int(bar.get_height()),
                ha="center",
                va="center",
                color="white",
                fontweight="bold",
                fontsize=9,
            )
    for bar in p2:
        if bar.get_height() > 0:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_y() + bar.get_height() / 2,
                int(bar.get_height()),
                ha="center",
                va="center",
                color="black",
                fontweight="bold",
                fontsize=9,
            )

    total = part + not_part
    apply_dynamic_ylim(ax, total)

    upper = ax.get_ylim()[1]

    for i in range(len(x)):
        tot = total.iloc[i]
        p = part.iloc[i]
        if tot > 0:
            pct = (p / tot) * 100
            ax.text(
                x[i],
                tot + (upper * 0.02),
                f"{pct:.0f}%",
                ha="center",
                va="bottom",
                fontweight="bold",
                fontsize=10,
            )

    ax.set_xticks(x)
    wrapped_labels = [
        textwrap.fill(str(label), width=15) for label in pivot["SCHOOL_NAME"]
    ]
    ax.set_xticklabels(wrapped_labels, rotation=0, ha="center")
    ax.set_xlabel("Schools", labelpad=15, fontweight="bold", fontsize=12)
    ax.set_ylabel("Students", labelpad=10, fontweight="bold", fontsize=12)

    ax.legend(loc="lower center", bbox_to_anchor=(0.5, 1.05), ncol=2)

    plt.tight_layout()

    image_path = os.path.join(IMAGE_FOLDER, "school_participation.png")
    plt.savefig(image_path, dpi=CHART_DPI)
    plt.close()

    return image_path


def generate_grade_overall_bar(df):

    summary = (
        df.groupby("GRADE")[["PARTICIPATED", "NOT_PARTICIPATED"]].sum().reset_index()
    )

    x = np.arange(len(summary))
    width = 0.6

    part = summary["PARTICIPATED"]
    not_part = summary["NOT_PARTICIPATED"]

    fig, ax = create_standard_chart((10, 5))

    c_part = get_theme_color("Participated")
    c_npart = get_theme_color("Not Participated")

    p1 = ax.bar(
        x,
        part,
        width,
        label="Participated",
        color=c_part,
        edgecolor="black",
        linewidth=1,
    )
    p2 = ax.bar(
        x,
        not_part,
        width,
        bottom=part,
        label="Not Participated",
        color=c_npart,
        edgecolor="black",
        linewidth=1,
    )
    total = part + not_part
    apply_dynamic_ylim(ax, total)
    upper = ax.get_ylim()[1]

    for bar in p1:
        if bar.get_height() > 0:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_y() + bar.get_height() / 2,
                int(bar.get_height()),
                ha="center",
                va="center",
                color="white",
                fontweight="bold",
                fontsize=9,
            )
    for bar in p2:
        if bar.get_height() > 0:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_y() + bar.get_height() / 2,
                int(bar.get_height()),
                ha="center",
                va="center",
                color="black",
                fontweight="bold",
                fontsize=9,
            )

    for i in range(len(x)):
        tot = total.iloc[i]
        p = part.iloc[i]
        if tot > 0:
            pct = (p / tot) * 100
            ax.text(
                x[i],
                tot + (upper * 0.02),
                f"{pct:.0f}%",
                ha="center",
                va="bottom",
                fontweight="bold",
                fontsize=10,
            )

    ax.set_xticks(x)
    ax.set_xticklabels(summary["GRADE"])
    ax.set_xlabel("Grades", labelpad=10, fontweight="bold", fontsize=12)
    ax.set_ylabel("Students", labelpad=10, fontweight="bold", fontsize=12)

    ax.legend(loc="lower center", bbox_to_anchor=(0.5, 1.05), ncol=2)

    plt.tight_layout()

    image_path = os.path.join(IMAGE_FOLDER, "grade_overall_participation.png")
    plt.savefig(image_path, dpi=CHART_DPI)
    plt.close()

    return image_path


def generate_school_grade_bar(df, school):

    sub = df[df["SCHOOL_NAME"] == school]

    x = np.arange(len(sub))
    width = 0.6

    part = sub["PARTICIPATED"]
    not_part = sub["NOT_PARTICIPATED"]

    fig, ax = create_standard_chart((10, 5))

    c_part = get_theme_color("Participated")
    c_npart = get_theme_color("Not Participated")

    p1 = ax.bar(
        x,
        part,
        width,
        label="Participated",
        color=c_part,
        edgecolor="black",
        linewidth=1,
    )
    p2 = ax.bar(
        x,
        not_part,
        width,
        bottom=part,
        label="Not Participated",
        color=c_npart,
        edgecolor="black",
        linewidth=1,
    )
    total = part + not_part
    apply_dynamic_ylim(ax, total)
    upper = ax.get_ylim()[1]

    for bar in p1:
        if bar.get_height() > 0:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_y() + bar.get_height() / 2,
                int(bar.get_height()),
                ha="center",
                va="center",
                color="white",
                fontweight="bold",
                fontsize=9,
            )
    for bar in p2:
        if bar.get_height() > 0:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_y() + bar.get_height() / 2,
                int(bar.get_height()),
                ha="center",
                va="center",
                color="black",
                fontweight="bold",
                fontsize=9,
            )

    for i in range(len(x)):
        tot = total.iloc[i]
        p = part.iloc[i]
        if tot > 0:
            pct = (p / tot) * 100
            ax.text(
                x[i],
                tot + (upper * 0.02),
                f"{pct:.0f}%",
                ha="center",
                va="bottom",
                fontweight="bold",
                fontsize=10,
            )

    ax.set_xticks(x)
    ax.set_xticklabels(sub["GRADE"])
    ax.set_xlabel("Grades", labelpad=10, fontweight="bold", fontsize=12)
    ax.set_ylabel("Students", labelpad=10, fontweight="bold", fontsize=12)

    ax.legend(loc="lower center", bbox_to_anchor=(0.5, 1.05), ncol=2)

    plt.tight_layout()

    safe_name = school.replace(" ", "_")
    image_path = os.path.join(IMAGE_FOLDER, f"{safe_name}_grade_participation.png")

    plt.savefig(image_path, dpi=CHART_DPI)
    plt.close()

    return image_path


# ==========================================================
# LOAD DATA
# ==========================================================

xls = pd.ExcelFile(INPUT_FILE)

mastery_df = load_sheet(xls, "MASTERY_COMPARISON_ALL")
qlo_df = load_sheet(xls, "QLO_PERFORMANCE_ALL")
school_performance_df = load_sheet(xls, "SCHOOL_PERFORMANCE_ALL")
participation_school_df = load_sheet(xls, "PARTICIPATION_SCHOOL_ALL")
participation_school_grade_df = load_sheet(xls, "PARTICIPATION_SCHOOL_GRADE_ALL")

for df in [participation_school_df, participation_school_grade_df]:
    if not df.empty and "TYPE" in df.columns:
        df["TYPE"] = (
            df["TYPE"].astype(str).str.strip().str.upper().str.replace(" ", "_")
        )

TEMPLATE_PATH = os.path.join(DATA_DIR, "template.pptx")
prs = Presentation(TEMPLATE_PATH)

# ==========================================================
# MASTERY SECTION
# ==========================================================

if not mastery_df.empty:
    for grade in sorted(mastery_df["GRADE"].dropna().unique()):
        grade_df = mastery_df[mastery_df["GRADE"] == grade]

        slide = duplicate_template_slide(prs, 0)
        enforce_template_title(slide, prs, f"{TITLES['MASTERY']} - Grade {grade}")

        image_path = generate_mastery_chart(grade_df, grade)

        add_centered_chart(slide, prs, image_path)

        txBox = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.1), Inches(4), Inches(0.3)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "Note: L is Mastery Level"
        p.font.size = Pt(11)
        p.font.italic = True
        p.font.color.rgb = RGBColor(80, 80, 80)

# ==========================================================
# QLO SECTION
# ==========================================================

for grade in sorted(qlo_df["GRADE"].unique()):
    grade_df = qlo_df[qlo_df["GRADE"] == grade]

    for subject in grade_df["SUBJECT"].unique():
        sub_df = grade_df[grade_df["SUBJECT"] == subject]

        slide = duplicate_template_slide(prs, 0)
        enforce_template_title(
            slide,
            prs,
            f"{TITLES['QLO']} - Grade {grade} - {SUBJECT_MAP.get(subject, subject)}",
        )

        image_path = generate_qlo_chart(sub_df, subject)
        add_centered_chart(slide, prs, image_path)

# ==========================================================
# SUBJECT PERFORMANCE SECTION
# ==========================================================

if not school_performance_df.empty:
    slide = duplicate_template_slide(prs, 0)
    enforce_template_title(slide, prs, "Overall Subject Performance (All Schools)")
    image_path = generate_overall_subject_performance(school_performance_df)
    add_centered_chart(slide, prs, image_path)

    for school in sorted(school_performance_df["SCHOOLNAME"].dropna().unique()):
        slide = duplicate_template_slide(prs, 0)
        enforce_template_title(
            slide, prs, f"{school} - Subject Performance (All Grades)"
        )
        image_path = generate_per_school_subject_performance(
            school_performance_df, school
        )
        add_centered_chart(slide, prs, image_path)

# ==========================================================
# PARTICIPATION SECTION
# ==========================================================

if not participation_school_df.empty:
    slide = duplicate_template_slide(prs, 0)
    enforce_template_title(slide, prs, "Overall Participation")
    image_path, total_regs = generate_overall_participation_pie(participation_school_df)

    box_width = Inches(3)
    box_height = Inches(0.5)
    reserve_margin = box_height + Inches(0.4)

    add_centered_chart(slide, prs, image_path, reserve_bottom=reserve_margin)

    left = (prs.slide_width - box_width) / 2
    bottom_safe_y = get_bottom_safe_y(slide, prs)
    top = bottom_safe_y - reserve_margin + Inches(0.2)

    txBox = slide.shapes.add_textbox(left, top, box_width, box_height)
    txBox.line.color.rgb = RGBColor(0, 0, 0)
    txBox.fill.background()
    tf = txBox.text_frame
    tf.text = f"Total Registrations: {total_regs}"
    for p in tf.paragraphs:
        p.font.size = Pt(14)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

    slide = duplicate_template_slide(prs, 0)
    enforce_template_title(slide, prs, "School-wise Participation")
    image_path = generate_school_participation_bar(participation_school_df)
    add_centered_chart(slide, prs, image_path)


if not participation_school_grade_df.empty:
    slide = duplicate_template_slide(prs, 0)
    enforce_template_title(slide, prs, "Grade-wise Overall Participation")
    image_path = generate_grade_overall_bar(participation_school_grade_df)
    add_centered_chart(slide, prs, image_path)

    for school in participation_school_grade_df["SCHOOL_NAME"].unique():
        slide = duplicate_template_slide(prs, 0)
        enforce_template_title(slide, prs, f"{school} - Grade-wise Participation")

        image_path = generate_school_grade_bar(participation_school_grade_df, school)

        add_centered_chart(slide, prs, image_path)

# ==========================================================
# SAVE FILE
# ==========================================================

rId = prs.slides._sldIdLst[0].rId
prs.part.drop_rel(rId)
del prs.slides._sldIdLst[0]
prs.save(OUTPUT_FILE)
print("ANALYSIS_REPORT.pptx created successfully.")
